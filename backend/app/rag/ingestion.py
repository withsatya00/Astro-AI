"""
Document ingestion pipeline for AstroAI Agent knowledge base.
Handles file validation, text extraction (PDF, DOCX, TXT, MD),
checksum calculation, deduplication, chunking, and embedding indexing.
"""
from typing import List, Dict, Any, Tuple, Optional
import os
import hashlib
from pathlib import Path
from sqlalchemy.orm import Session
from backend.app.models.models import Document, DocumentChunk
from backend.app.rag.chunking import chunk_document_text
from backend.app.rag.embeddings import EmbeddingService
from backend.app.core.logging import logger

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB

def compute_sha256(file_path: str) -> str:
    """Computes SHA256 hash of a file."""
    sha = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            sha.update(chunk)
    return sha.hexdigest()

def extract_text_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """Extracts text page-by-page from PDF using PyMuPDF (fitz) or pypdf."""
    pages = []
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            pages.append({
                "page_number": page_num + 1,
                "text": text,
                "section": f"Page {page_num + 1}"
            })
        doc.close()
        return pages
    except Exception as e:
        logger.warning(f"PyMuPDF failed on {file_path} ({e}), trying pypdf fallback.")

    # Fallback to pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append({
                "page_number": idx + 1,
                "text": text,
                "section": f"Page {idx + 1}"
            })
        return pages
    except Exception as e:
        logger.error(f"Failed to extract PDF text from {file_path}: {e}")
        raise ValueError(f"Unable to parse PDF: {e}")

def extract_text_from_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extracts paragraphs and headings from DOCX file."""
    try:
        import docx
        doc = docx.Document(file_path)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text)
        combined = "\n\n".join(full_text)
        return [{"page_number": 1, "text": combined, "section": "Document Content"}]
    except Exception as e:
        raise ValueError(f"Unable to parse DOCX: {e}")

def extract_text_from_plain_text(file_path: str) -> List[Dict[str, Any]]:
    """Extracts plain text or markdown file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return [{"page_number": 1, "text": content, "section": "Text Document"}]
    except Exception as e:
        raise ValueError(f"Unable to read text file: {e}")

class IngestionPipeline:
    @staticmethod
    def process_and_index_file(
        db: Session,
        file_path: str,
        title: str,
        author: Optional[str] = None,
        publisher: Optional[str] = None,
        license_str: str = "Public Domain / CC",
        source_url: Optional[str] = None,
        permission_status: str = "PERMITTED"
    ) -> Document:
        """
        Validates, extracts, chunks, generates embeddings, and indexes a file into the database.
        """
        path_obj = Path(file_path)
        if not path_obj.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_size = path_obj.stat().st_size
        if file_size > MAX_FILE_SIZE_BYTES:
            raise ValueError(f"File size ({file_size} bytes) exceeds 25MB limit.")

        ext = path_obj.suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            raise ValueError(f"Unsupported file format: {ext}. Allowed formats: {ALLOWED_EXTENSIONS}")

        checksum = compute_sha256(file_path)

        # Duplicate check
        existing_doc = db.query(Document).filter(Document.checksum == checksum).first()
        if existing_doc:
            logger.info(f"Document with checksum {checksum} already indexed as '{existing_doc.title}'.")
            return existing_doc

        # Create Document record
        doc_record = Document(
            title=title,
            author=author,
            publisher=publisher,
            source_url=source_url,
            license=license_str,
            permission_status=permission_status,
            file_path=str(file_path),
            file_type=ext.lstrip("."),
            file_size_bytes=file_size,
            checksum=checksum,
            status="PROCESSING"
        )
        db.add(doc_record)
        db.commit()
        db.refresh(doc_record)

        try:
            # 1. Extract text
            if ext == ".pdf":
                pages = extract_text_from_pdf(file_path)
            elif ext == ".docx":
                pages = extract_text_from_docx(file_path)
            else:
                pages = extract_text_from_plain_text(file_path)

            # 2. Chunk text
            chunks_data = chunk_document_text(pages, chunk_size=650, chunk_overlap=120)

            # 3. Generate embeddings and create Chunk records
            chunk_records = []
            for item in chunks_data:
                vec = EmbeddingService.get_embedding(item["content"])
                chunk_obj = DocumentChunk(
                    document_id=doc_record.id,
                    chunk_index=item["chunk_index"],
                    content=item["content"],
                    page_number=item.get("page_number", 1),
                    section_title=item.get("section_title"),
                    embedding=vec
                )
                chunk_records.append(chunk_obj)

            db.bulk_save_objects(chunk_records)
            doc_record.status = "INDEXED"
            doc_record.chunk_count = len(chunk_records)
            db.commit()
            db.refresh(doc_record)

            logger.info(f"Successfully indexed document '{title}' ({doc_record.id}) with {len(chunk_records)} chunks.")
            return doc_record

        except Exception as e:
            db.rollback()
            doc_record.status = "FAILED"
            doc_record.error_details = str(e)
            db.commit()
            logger.error(f"Failed to ingest document '{title}': {e}")
            raise e

    @staticmethod
    def delete_document(db: Session, doc_id: str) -> bool:
        """Deletes a document and all its associated indexed chunks."""
        doc = db.query(Document).filter(Document.id == doc_id).first()
        if not doc:
            return False
        db.delete(doc)
        db.commit()
        logger.info(f"Deleted document {doc_id} and its associated chunks.")
        return True
